"""Gravação do quadro de triagem em planilha ou CSV, sempre na máquina local."""

from __future__ import annotations

import csv
import json
import re
import time
from pathlib import Path


def _colunas(linhas: list[dict]) -> list[str]:
    ordem: list[str] = []
    for linha in linhas:
        for chave in linha:
            if chave not in ordem:
                ordem.append(chave)
    return ordem


def _celula(valor) -> str | int | float:
    if isinstance(valor, (list, tuple)):
        return "; ".join(str(v) for v in valor)
    if isinstance(valor, dict):
        return json.dumps(valor, ensure_ascii=False)
    if valor is None:
        return ""
    return valor


def gravar(linhas: list[dict], pasta: str, nome: str | None, formato: str) -> Path:
    if not linhas:
        raise ValueError("Nenhuma linha para exportar.")
    formato = formato.lower()
    if formato not in {"xlsx", "csv"}:
        raise ValueError("Formato deve ser xlsx ou csv.")
    destino_dir = Path(pasta).expanduser()
    destino_dir.mkdir(parents=True, exist_ok=True)
    base = re.sub(r"[^\w\-]+", "_", nome or "") or f"triagem-{time.strftime('%Y%m%d-%H%M%S')}"
    destino = destino_dir / f"{base}.{formato}"
    colunas = _colunas(linhas)

    if formato == "csv":
        with open(destino, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(colunas)
            for linha in linhas:
                w.writerow([_celula(linha.get(c)) for c in colunas])
        return destino

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Triagem"
    ws.append(colunas)
    for linha in linhas:
        ws.append([_celula(linha.get(c)) for c in colunas])
    cabecalho = PatternFill("solid", fgColor="1F3A5F")
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = cabecalho
        cel.alignment = Alignment(vertical="center", wrap_text=True)
    for i, col in enumerate(colunas, start=1):
        maior = max(len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, maior + 2), 60)
    for linha in ws.iter_rows(min_row=2):
        for cel in linha:
            cel.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(destino)
    return destino
